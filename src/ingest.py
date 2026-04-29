from __future__ import annotations

import base64
import csv
import io
import json
import mimetypes
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from uuid import uuid4

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageSequence
from pptx import Presentation
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text

from .config import AppConfig


DOCUMENT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".rtf",
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".odt",
}

IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".gif",
    ".bmp",
    ".tif",
    ".tiff",
}

ZIP_EXTENSIONS = {".zip"}
SUPPORTED_EXTENSIONS = sorted(DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS | ZIP_EXTENSIONS)
MAX_NESTED_ZIP_DEPTH = 8
IngestProgressCallback = Callable[[str], None]
IngestProgressEventCallback = Callable[[str, int, int], None]


@dataclass
class IngestedItem:
    name: str
    kind: str
    extension: str
    stored_path: str | None
    text: str
    image_data_url: str | None = None
    mime_type: str | None = None
    error: str | None = None

    @property
    def char_count(self) -> int:
        return len(self.text or "")


def sanitize_filename(name: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9._ -]+", "_", name).strip(" .")
    return clean[:160] or f"upload-{uuid4().hex}"


def _safe_zip_name(name: str, is_dir: bool = False) -> str | None:
    normalized = name.replace("\\", "/")
    path = Path(normalized)
    if path.is_absolute() or ".." in path.parts:
        return None
    if is_dir or normalized.endswith("/"):
        return None
    return str(path).replace("\\", "/")


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _truncate(text: str, limit: int) -> str:
    text = text.strip()
    if len(text) <= limit:
        return text
    return text[:limit] + "\n\n[Truncated for context budget.]"


def _emit_progress(
    callback: IngestProgressEventCallback | None,
    message: str,
    completed: int,
    total: int,
) -> None:
    if callback:
        callback(message, max(0, completed), max(1, total))


def _parse_pdf(
    data: bytes,
    *,
    max_chars: int,
    display_name: str,
    progress_event_callback: IngestProgressEventCallback | None = None,
) -> str:
    reader = PdfReader(io.BytesIO(data), strict=False)
    if reader.is_encrypted:
        decrypt_result = reader.decrypt("")
        if decrypt_result == 0:
            raise ValueError("Encrypted PDF requires a password.")

    total_pages = len(reader.pages)
    _emit_progress(progress_event_callback, f"Reading PDF: {display_name} (0/{total_pages} pages)", 0, total_pages)
    pages = []
    extracted_chars = 0
    for index, page in enumerate(reader.pages, start=1):
        _emit_progress(
            progress_event_callback,
            f"Reading PDF: {display_name} ({index}/{total_pages} pages)",
            index - 1,
            total_pages,
        )
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            pages.append(f"[Page {index}]\n[Could not extract text from this page: {exc}]")
            _emit_progress(
                progress_event_callback,
                f"Skipped unreadable PDF page {index}/{total_pages}: {display_name}",
                index,
                total_pages,
            )
            continue
        if text.strip():
            page_text = f"[Page {index}]\n{text.strip()}"
            pages.append(page_text)
            extracted_chars += len(page_text)
        _emit_progress(
            progress_event_callback,
            f"Read PDF page {index}/{total_pages}: {display_name}",
            index,
            total_pages,
        )
        if extracted_chars >= max_chars:
            pages.append("[PDF text truncated after reaching the per-file context limit.]")
            break
    return "\n\n".join(pages)


def _parse_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    chunks: list[str] = []
    chunks.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def _parse_pptx(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    chunks: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_text: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())
        if slide_text:
            chunks.append(f"[Slide {slide_index}]\n" + "\n".join(slide_text))
    return "\n\n".join(chunks)


def _parse_xlsx(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks: list[str] = []
    max_rows_per_sheet = 400
    max_cols_per_sheet = 30
    for sheet in workbook.worksheets:
        rows: list[str] = []
        for row_index, row in enumerate(
            sheet.iter_rows(max_row=max_rows_per_sheet, max_col=max_cols_per_sheet, values_only=True),
            start=1,
        ):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                rows.append("\t".join(values).rstrip())
            if row_index >= max_rows_per_sheet:
                break
        if rows:
            chunks.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
    return "\n\n".join(chunks)


def _parse_csv(data: bytes) -> str:
    text = _decode_text(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel
    rows: list[str] = []
    reader = csv.reader(io.StringIO(text), dialect=dialect)
    for index, row in enumerate(reader, start=1):
        rows.append("\t".join(row))
        if index >= 1000:
            rows.append("[Truncated CSV after 1000 rows.]")
            break
    return "\n".join(rows)


def _parse_odt(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        content = archive.read("content.xml")
    soup = BeautifulSoup(content, "xml")
    return "\n".join(node.get_text(" ", strip=True) for node in soup.find_all("p") if node.get_text(strip=True))


def _parse_html(data: bytes) -> str:
    soup = BeautifulSoup(_decode_text(data), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n", strip=True)


def _image_to_data_url(data: bytes, extension: str) -> tuple[str, str]:
    image = Image.open(io.BytesIO(data))
    if extension == ".gif":
        frames = [frame.copy() for frame in ImageSequence.Iterator(image)]
        if len(frames) > 1:
            image = frames[0]
    image.thumbnail((2000, 2000))

    target_format = "PNG"
    mime_type = "image/png"
    if extension in {".jpg", ".jpeg"}:
        target_format = "JPEG"
        mime_type = "image/jpeg"
    elif extension == ".webp":
        target_format = "WEBP"
        mime_type = "image/webp"
    elif extension == ".gif" and getattr(image, "is_animated", False) is False:
        target_format = "GIF"
        mime_type = "image/gif"

    output = io.BytesIO()
    if target_format == "JPEG" and image.mode not in {"RGB", "L"}:
        image = image.convert("RGB")
    image.save(output, format=target_format)
    encoded = base64.b64encode(output.getvalue()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}", mime_type


def _parse_file(
    data: bytes,
    display_name: str,
    stored_path: str | None,
    config: AppConfig,
    progress_event_callback: IngestProgressEventCallback | None = None,
) -> IngestedItem:
    extension = Path(display_name).suffix.lower()
    try:
        if extension in IMAGE_EXTENSIONS:
            data_url, mime_type = _image_to_data_url(data, extension)
            return IngestedItem(display_name, "image", extension, stored_path, "", data_url, mime_type)

        if extension in {".txt", ".md", ".markdown", ".xml"}:
            text = _decode_text(data)
        elif extension == ".json":
            try:
                text = json.dumps(json.loads(_decode_text(data)), indent=2, ensure_ascii=False)
            except json.JSONDecodeError:
                text = _decode_text(data)
        elif extension == ".csv":
            text = _parse_csv(data)
        elif extension in {".html", ".htm"}:
            text = _parse_html(data)
        elif extension == ".rtf":
            text = rtf_to_text(_decode_text(data))
        elif extension == ".pdf":
            text = _parse_pdf(
                data,
                max_chars=config.max_text_chars_per_file,
                display_name=display_name,
                progress_event_callback=progress_event_callback,
            )
        elif extension == ".docx":
            text = _parse_docx(data)
        elif extension == ".pptx":
            text = _parse_pptx(data)
        elif extension == ".xlsx":
            text = _parse_xlsx(data)
        elif extension == ".odt":
            text = _parse_odt(data)
        else:
            return IngestedItem(display_name, "unsupported", extension, stored_path, "", error="Unsupported file type.")
        return IngestedItem(
            display_name,
            "document",
            extension,
            stored_path,
            _truncate(text, config.max_text_chars_per_file),
            mime_type=mimetypes.guess_type(display_name)[0],
        )
    except Exception as exc:
        return IngestedItem(display_name, "error", extension, stored_path, "", error=f"Could not read file: {exc}")


def _write_upload(data: bytes, upload_root: Path, user_id: int, purpose: str, display_name: str) -> str:
    folder = upload_root / str(user_id) / purpose
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{uuid4().hex}-{sanitize_filename(display_name)}"
    path.write_bytes(data)
    return str(path)


def ingest_uploads(
    uploaded_files: Iterable,
    config: AppConfig,
    user_id: int,
    purpose: str,
    progress_callback: IngestProgressCallback | None = None,
    progress_event_callback: IngestProgressEventCallback | None = None,
) -> list[IngestedItem]:
    items: list[IngestedItem] = []
    uploaded_list = list(uploaded_files)
    total_uploads = max(1, len(uploaded_list))
    for upload_index, uploaded in enumerate(uploaded_list, start=1):
        name = getattr(uploaded, "name", "upload")
        if progress_callback:
            progress_callback(f"Reading upload: {name}")
        _emit_progress(
            progress_event_callback,
            f"Reading upload {upload_index}/{total_uploads}: {name}",
            upload_index - 1,
            total_uploads,
        )
        data = uploaded.getvalue()
        stored_path = _write_upload(data, config.upload_dir, user_id, purpose, name)
        extension = Path(name).suffix.lower()
        if extension in ZIP_EXTENSIONS:
            items.extend(
                _ingest_zip(
                    data,
                    archive_name=name,
                    archive_path=stored_path,
                    config=config,
                    user_id=user_id,
                    purpose=purpose,
                    progress_callback=progress_callback,
                    progress_event_callback=progress_event_callback,
                )
            )
        else:
            items.append(_parse_file(data, name, stored_path, config, progress_event_callback))
        _emit_progress(
            progress_event_callback,
            f"Finished upload {upload_index}/{total_uploads}: {name}",
            upload_index,
            total_uploads,
        )
    return items


def _ingest_zip(
    data: bytes,
    archive_name: str,
    archive_path: str,
    config: AppConfig,
    user_id: int,
    purpose: str,
    depth: int = 0,
    progress_callback: IngestProgressCallback | None = None,
    progress_event_callback: IngestProgressEventCallback | None = None,
) -> list[IngestedItem]:
    items: list[IngestedItem] = []
    if depth > MAX_NESTED_ZIP_DEPTH:
        return [
            IngestedItem(
                archive_name,
                "error",
                ".zip",
                archive_path,
                "",
                error="Nested zip depth exceeded.",
            )
        ]
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            total_entries = max(1, len(entries))
            for entry_index, info in enumerate(entries, start=1):
                safe_name = _safe_zip_name(info.filename, info.is_dir())
                if safe_name is None:
                    continue
                extension = Path(safe_name).suffix.lower()
                nested_name = f"{archive_name}/{safe_name}"
                if progress_callback:
                    progress_callback(f"Unzipping: {nested_name}")
                _emit_progress(
                    progress_event_callback,
                    f"Unzipping {entry_index}/{total_entries}: {nested_name}",
                    entry_index - 1,
                    total_entries,
                )
                try:
                    with archive.open(info) as file_handle:
                        file_data = file_handle.read()
                except Exception as exc:
                    items.append(
                        IngestedItem(
                            nested_name,
                            "error",
                            extension,
                            archive_path,
                            "",
                            error=f"Could not extract file from zip: {exc}",
                        )
                    )
                    continue

                if extension in ZIP_EXTENSIONS:
                    nested_path = _write_upload(file_data, config.upload_dir, user_id, purpose, nested_name)
                    items.extend(
                        _ingest_zip(
                            file_data,
                            archive_name=nested_name,
                            archive_path=nested_path,
                            config=config,
                            user_id=user_id,
                            purpose=purpose,
                            depth=depth + 1,
                            progress_callback=progress_callback,
                            progress_event_callback=progress_event_callback,
                        )
                    )
                    continue

                if extension not in DOCUMENT_EXTENSIONS and extension not in IMAGE_EXTENSIONS:
                    items.append(
                        IngestedItem(
                            nested_name,
                            "unsupported",
                            extension,
                            archive_path,
                            "",
                            error="Unsupported file type inside zip.",
                        )
                    )
                    continue
                nested_path = _write_upload(file_data, config.upload_dir, user_id, purpose, nested_name)
                items.append(_parse_file(file_data, nested_name, nested_path, config, progress_event_callback))
                _emit_progress(
                    progress_event_callback,
                    f"Finished zip entry {entry_index}/{total_entries}: {nested_name}",
                    entry_index,
                    total_entries,
                )
    except zipfile.BadZipFile:
        items.append(
            IngestedItem(
                archive_name,
                "error",
                ".zip",
                archive_path,
                "",
                error="Invalid zip file.",
            )
        )
    return items


def build_manifest(items: Iterable[IngestedItem]) -> list[dict[str, object]]:
    manifest = []
    for item in items:
        manifest.append(
            {
                "name": item.name,
                "kind": item.kind,
                "extension": item.extension,
                "chars": item.char_count,
                "has_image": bool(item.image_data_url),
                "error": item.error,
            }
        )
    return manifest


def combine_text_context(items: Iterable[IngestedItem], max_chars: int) -> str:
    chunks: list[str] = []
    remaining = max_chars
    for item in items:
        if item.error or not item.text.strip():
            continue
        header = f"\n\n--- SOURCE: {item.name} ---\n"
        block = header + item.text.strip()
        if len(block) > remaining:
            chunks.append(block[:remaining] + "\n[Context truncated.]")
            break
        chunks.append(block)
        remaining -= len(block)
        if remaining <= 0:
            break
    return "".join(chunks).strip()


def collect_images(items: Iterable[IngestedItem], max_images: int) -> list[tuple[str, str]]:
    images = []
    for item in items:
        if item.image_data_url:
            images.append((item.name, item.image_data_url))
            if len(images) >= max_images:
                break
    return images
